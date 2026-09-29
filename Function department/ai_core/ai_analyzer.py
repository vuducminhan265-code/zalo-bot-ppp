import os
import json
import re
import sys
import shutil
import tempfile
import time
from dotenv import load_dotenv
from google import genai
from google.genai import types
import pypdf

import base64
load_dotenv()

DEFAULT_GEMINI_KEY = base64.b64decode("QVEuQWI4Uk42THE0UVF6NzRkbmNHa0d6SzVmYVBfdl9RSDRSMGphTnVtYi1ibkVSS3JNZUE=").decode()
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY") or DEFAULT_GEMINI_KEY

class TaskAIAnalyzer:
    def __init__(self, api_key: str = None):
        self.api_key = api_key or GEMINI_API_KEY or DEFAULT_GEMINI_KEY
        if self.api_key:
            self.client = genai.Client(api_key=self.api_key)
        else:
            self.client = None

    def extract_task_number_from_filename(self, filename: str) -> str:
        """Trích xuất số Task ID đứng đầu tên file (Ví dụ: '1. 23920 STC-PPPQLN.pdf' -> '1')"""
        base_name = os.path.basename(filename)
        match = re.match(r"^(\d+)[.\s\-_]", base_name)
        if match:
            return match.group(1)
        
        match_task = re.search(r"(?:task|nhiem\s*vu|nv)[.\s\-_]*(\d+)", base_name, re.IGNORECASE)
        if match_task:
            return match_task.group(1)
            
        return ""

    def _extract_text_fallback(self, file_path: str) -> str:
        """Trích xuất nội dung văn bản trực tiếp từ file docx, doc, pdf nếu cần"""
        ext = os.path.splitext(file_path)[1].lower()
        if ext == ".pdf":
            try:
                reader = pypdf.PdfReader(file_path)
                pages_text = [p.extract_text() or "" for p in reader.pages[:10]]
                return "\n".join(pages_text)
            except Exception:
                pass
        elif ext in [".docx", ".doc"]:
            try:
                import docx
                doc = docx.Document(file_path)
                return "\n".join([p.text for p in doc.paragraphs if p.text])
            except Exception:
                pass
            try:
                with open(file_path, "rb") as f:
                    raw = f.read()
                text_utf16 = raw.decode("utf-16le", errors="ignore")
                matches = re.findall(r"[\w\s,.:;/–—\(\)\-]{10,}", text_utf16)
                if matches:
                    return " ".join(matches)
            except Exception:
                pass
        elif ext in [".xlsx", ".xls"]:
            try:
                import openpyxl
                wb = openpyxl.load_workbook(file_path, data_only=True)
                lines = []
                for sheet in wb.worksheets[:3]:
                    for row in sheet.iter_rows(values_only=True):
                        row_vals = [str(val) for val in row if val is not None]
                        if row_vals:
                            lines.append(" | ".join(row_vals))
                return "\n".join(lines[:200])
            except Exception:
                pass
        return ""

    def review_submission(self, file_path: str, assignee_name: str, pending_tasks: list) -> dict:
        """
        Đọc và phân tích file nộp:
        1. Bóc tách thể thức hành chính: Số ký hiệu, Ngày tháng, Trích yếu nội dung văn bản.
        2. Đối chiếu số Task ID từ tên file hoặc nội dung.
        3. Kiểm tra tính nhất quán giữa Trích yếu trong file và Trích yếu quy định trên Google Sheet.
        4. Kết luận:
           - 'completed': Đúng trích yếu, đúng sản phẩm -> Hoàn thành.
           - 'needs_revision': Không đúng trích yếu / sai lệch nội dung -> Đang sửa lại.
        """
        filename = os.path.basename(file_path)
        hint_task_id = self.extract_task_number_from_filename(filename)
        local_text = self._extract_text_fallback(file_path)

        prompt = f"""
Bạn là Trợ lý Nghiệm thu Công việc & Thẩm định Thể thức Văn bản Hành chính (Bộ phận PPP & Quản lý Nợ - Sở Tài Chính TP.HCM).
Thành viên: {assignee_name} vừa nộp tệp: "{filename}".
Gợi ý Task ID từ tên file: {hint_task_id if hint_task_id else "Không có số đứng đầu (cần phân tích theo nội dung văn bản)"}.

DANH SÁCH NHIỆM VỤ ĐANG CHỜ HOÀN THÀNH CỦA CHUYÊN VIÊN TRÊN GOOGLE SHEETS:
{json.dumps(pending_tasks, ensure_ascii=False, indent=2)}

NỘI DUNG VĂN BẢN TRÍCH XUẤT TỪ FILE:
\"\"\"
{local_text[:12000] if local_text else "Vui lòng phân tích tệp đính kèm trực tiếp"}
\"\"\"

NHIỆM VỤ PHÂN TÍCH:
1. BÓC TÁCH THỂ THỨC VĂN BẢN:
   - Số ký hiệu văn bản (Ví dụ: "23920/STC-PPP&QLN", "145/TTr-STC" hoặc rỗng nếu không có).
   - Ngày tháng năm ban hành văn bản (Định dạng "DD/MM/YYYY" hoặc chuỗi ngày tháng).
   - Trích yếu thực tế trong văn bản (Nội dung chính tóm tắt văn bản nói về việc gì).

2. ĐỐI SOÁT VỚI DANH SÁCH NHIỆM VỤ:
   - Xác định file này thuộc Task ID nào trong danh sách nhiệm vụ trên.
   - So sánh TRÍCH YẾU TRONG FILE với TRÍCH YẾU / TÊN CÔNG VIỆC QUY ĐỊNH của Task đó:
     * TRƯỜNG HỢP 1 (HỢP LỆ): Trích yếu và nội dung file KHỚP ĐÚNG với nhiệm vụ -> Trạng thái: "completed" (Đã hoàn thành).
     * TRƯỜNG HỢP 2 (KHÔNG KHỚP): Số task ghi đúng nhưng nội dung trích yếu KHÔNG KHỚP hoặc sai lệch -> Trạng thái: "needs_revision" (Đang sửa lại).

3. TỔNG KẾT & ĐÁNH GIÁ:
   - Viết 1-2 câu nhận xét thẩm định ngắn gọn, súc tích.

BẮT BUỘC TRẢ VỀ ĐỊNH DẠNG JSON DUY NHẤT (không bọc text ngoài JSON):
{{
  "task_id": "Mã số Task ID (Ví dụ: 1, 2, 3...)",
  "status_decision": "completed" hoặc "needs_revision",
  "document_number": "Số ký hiệu văn bản hoặc rỗng",
  "document_date": "Ngày tháng ban hành văn bản hoặc rỗng",
  "actual_summary": "Trích yếu thực tế bóc tách được từ trong file",
  "ai_review_comment": "Nhận xét thẩm định của AI (ngắn gọn 1-2 câu)",
  "discrepancy_reason": "Lý do sai lệch nếu có (hoặc null)",
  "resolution_options": [
    "Lựa chọn 1: Cập nhật lại Trích yếu trên Google Sheets cho khớp với văn bản thực tế",
    "Lựa chọn 2: Chỉnh sửa và nộp lại đúng văn bản theo yêu cầu của nhiệm vụ"
  ]
}}
"""

        # Thử các model Gemini theo thứ tự ưu tiên (Tối ưu Lite Model, tiết kiệm Token)
        models_to_try = ["gemini-3.5-flash-lite", "gemini-3.1-flash-lite", "gemini-flash-lite-latest", "gemini-3.5-flash"]
        
        last_error = None
        for model_name in models_to_try:
            for attempt in range(2):
                try:
                    # Nếu có text trích xuất trực tiếp, gửi kèm prompt
                    contents = [prompt]
                    
                    # Nếu không có text hoặc muốn gửi cả file
                    if not local_text:
                        ext = os.path.splitext(file_path)[1].lower()
                        temp_dir = tempfile.gettempdir()
                        ascii_temp = os.path.join(temp_dir, f"temp_upload_{os.getpid()}_{attempt}{ext}")
                        shutil.copyfile(file_path, ascii_temp)
                        try:
                            uploaded_file = self.client.files.upload(file=ascii_temp)
                            contents.insert(0, uploaded_file)
                        finally:
                            if os.path.exists(ascii_temp):
                                os.remove(ascii_temp)
                    
                    resp = self.client.models.generate_content(
                        model=model_name,
                        contents=contents,
                        config=types.GenerateContentConfig(
                            response_mime_type="application/json",
                            temperature=0.1
                        )
                    )
                    
                    raw_text = resp.text.strip()
                    raw_text = re.sub(r"^```json\s*", "", raw_text)
                    raw_text = re.sub(r"```$", "", raw_text).strip()
                    
                    data = json.loads(raw_text)
                    data["success"] = True
                    return data
                except Exception as e:
                    last_error = str(e)
                    time.sleep(1.5)

        return {
            "success": False,
            "error": f"Lỗi phân tích AI sau các lần thử: {last_error}"
        }
